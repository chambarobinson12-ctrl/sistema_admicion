from datetime import date
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.views.decorators.clickjacking import xframe_options_sameorigin
from django.contrib import messages
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q, Count
from django.utils import timezone
from django.http import HttpResponse
from django.utils.http import url_has_allowed_host_and_scheme

from convocatorias.models import Convocatoria, Carrera, CupoCarrera
from postulantes.models import Inscripcion, Documento, ConfiguracionProceso
from evaluacion.models import Examen, Resultado
from .decorators import panel_required
from .paginacion import paginar
from .notificaciones import notificar_documento, notificar_decision_inscripcion, notificar_reintegro, notificar_comentario
from .forms import CarreraForm, ConvocatoriaForm, CupoCarreraFormSet, ExamenForm, ConfiguracionProcesoForm


def _convocatoria_activa():
    return Convocatoria.objects.filter(activa=True).order_by('fecha_fin').first()
def _convocatoria_seleccionada(request):
    """Devuelve la convocatoria elegida en el filtro ?proceso=,
    o la activa si no se eligió ninguna."""
    proceso_id = request.GET.get('proceso', '')
    convocatoria = None
    if proceso_id.isdigit():
        convocatoria = Convocatoria.objects.filter(pk=proceso_id).first()
    return convocatoria or _convocatoria_activa()

def _inscripciones_filtradas(request, convocatoria):
    """Aplica a la lista de inscripciones los mismos filtros (q, carrera,
    estado) que usa el listado de Postulantes del panel, para que la lista
    en pantalla y el Excel exportado siempre muestren exactamente lo mismo."""
    inscripciones = Inscripcion.objects.select_related('postulante', 'postulante__usuario', 'carrera', 'convocatoria').prefetch_related('documentos')
    if convocatoria:
        inscripciones = inscripciones.filter(convocatoria=convocatoria)

    q = request.GET.get('q', '').strip()
    carrera_id = request.GET.get('carrera', '')
    estado = request.GET.get('estado', '')

    if q:
        inscripciones = inscripciones.filter(
            Q(postulante__nombres__icontains=q)
            | Q(postulante__apellidos__icontains=q)
            | Q(postulante__usuario__cedula__icontains=q)
            | Q(numero_postulacion__icontains=q)
        )
    if carrera_id:
        inscripciones = inscripciones.filter(carrera_id=carrera_id)
    if estado:
        inscripciones = inscripciones.filter(estado=estado)

    return inscripciones, q, carrera_id, estado


@panel_required
def dashboard(request):
    convocatoria = _convocatoria_seleccionada(request)

    inscripciones, q, carrera_id, estado = _inscripciones_filtradas(request, convocatoria)

    # Avance de cada postulante: completo (subió todos los certificados) o
    # incompleto (le falta alguno o dejó alguno en blanco).
    inscripciones = list(inscripciones.prefetch_related('documentos'))
    total_postulantes = len(inscripciones)
    procesos_completos = sum(1 for i in inscripciones if i.proceso_completo)
    procesos_incompletos = total_postulantes - procesos_completos
    filtro_avance = request.GET.get('avance', '')
    if filtro_avance == 'completo':
        inscripciones = [i for i in inscripciones if i.proceso_completo]
    elif filtro_avance == 'incompleto':
        inscripciones = [i for i in inscripciones if not i.proceso_completo]

    cupos_disponibles = 0
    cupos_totales = 0
    if convocatoria:
        for cupo in convocatoria.cupocarrera_set.all():
            cupos_totales += cupo.cupos
            aprobados = Inscripcion.objects.filter(
                convocatoria=convocatoria, carrera=cupo.carrera, estado='aprobada'
            ).count()
            cupos_disponibles += max(cupo.cupos - aprobados, 0)

    pagina = paginar(request, inscripciones)

    carreras = Carrera.objects.filter(cupocarrera__convocatoria=convocatoria).distinct() if convocatoria else Carrera.objects.none()

    return render(request, 'panel/dashboard.html', {
        'convocatoria': convocatoria,
        'pagina': pagina,
        'total_postulantes': total_postulantes,
        'procesos_completos': procesos_completos,
        'procesos_incompletos': procesos_incompletos,
        'filtro_avance': filtro_avance,
        'cupos_disponibles': cupos_disponibles,
        'cupos_totales': cupos_totales,
        'carreras': carreras,
        'filtro_q': q,
        'filtro_carrera': carrera_id,
        'filtro_estado': estado,
        'estados': Inscripcion.ESTADO_CHOICES,
        'convocatorias': Convocatoria.objects.order_by('-fecha_inicio'),
        'filtro_proceso': str(convocatoria.id) if convocatoria else '',
    })


@panel_required
def exportar_postulantes_excel(request):
    """Genera un .xlsx con la lista de postulantes (respetando los mismos
    filtros de búsqueda/carrera/estado del listado), con un diseño similar
    al catálogo institucional: encabezado verde con el nombre del instituto
    y una fila de títulos de columna en verde claro."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    convocatoria = _convocatoria_seleccionada(request)
    inscripciones, q, carrera_id, estado = _inscripciones_filtradas(request, convocatoria)
    inscripciones = list(inscripciones.order_by('postulante__apellidos', 'postulante__nombres'))
    filtro_avance = request.GET.get('avance', '')
    if filtro_avance == 'completo':
        inscripciones = [i for i in inscripciones if i.proceso_completo]
    elif filtro_avance == 'incompleto':
        inscripciones = [i for i in inscripciones if not i.proceso_completo]

    columnas = [
        ('N°', 6),
        ('Postulante', 30),
        ('Cédula', 16),
        ('Contacto', 30),
        ('Carrera', 34),
        ('N.° de postulación', 20),
        ('Proceso de admisión', 26),
        ('Fecha de inscripción', 18),
        ('Proceso', 20),
        ('Documentos subidos', 20),
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Postulantes'

    verde_fuerte = 'FF8BC34A'
    verde_claro = 'FFE8F5E9'
    blanco = 'FFFFFFFF'
    negro = 'FF000000'
    borde_fino = Side(style='thin', color='FFBDBDBD')
    borde_celda = Border(left=borde_fino, right=borde_fino, top=borde_fino, bottom=borde_fino)

    ultima_col = len(columnas)
    ultima_letra = get_column_letter(ultima_col)

    ws.merge_cells(f'A1:{ultima_letra}2')
    titulo = ws['A1']
    titulo.value = 'INSTITUTO SUPERIOR TECNOLÓGICO AMAZÓNICO'
    titulo.font = Font(name='Calibri', size=14, bold=True, color=blanco)
    titulo.fill = PatternFill('solid', fgColor=verde_fuerte)
    titulo.alignment = Alignment(horizontal='center', vertical='center')
    for col in range(1, ultima_col + 1):
        ws.cell(row=1, column=col).fill = PatternFill('solid', fgColor=verde_fuerte)
        ws.cell(row=2, column=col).fill = PatternFill('solid', fgColor=verde_fuerte)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 20

    subtitulo_partes = ['Listado de postulantes']
    if convocatoria:
        subtitulo_partes.append(convocatoria.nombre)
    subtitulo_partes.append(f'Generado el {timezone.localtime().strftime("%d/%m/%Y %H:%M")}')
    ws.merge_cells(f'A3:{ultima_letra}3')
    subtitulo = ws['A3']
    subtitulo.value = ' · '.join(subtitulo_partes)
    subtitulo.font = Font(name='Calibri', size=10, italic=True, color='FF33691E')
    subtitulo.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[3].height = 18

    fila_encabezado = 4
    for col, (etiqueta, ancho) in enumerate(columnas, start=1):
        celda = ws.cell(row=fila_encabezado, column=col, value=etiqueta)
        celda.font = Font(name='Calibri', size=11, bold=True, color=negro)
        celda.fill = PatternFill('solid', fgColor=verde_claro)
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        celda.border = borde_celda
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.row_dimensions[fila_encabezado].height = 26

    fila = fila_encabezado + 1
    for numero, inscripcion in enumerate(inscripciones, start=1):
        docs_subidos = inscripcion.documentos_subidos_total
        docs_totales = len(Documento.TIPO_CHOICES)
        valores = [
            numero,
            inscripcion.postulante.nombre_completo,
            inscripcion.postulante.usuario.cedula or '—',
            # Teléfono; si no lo registró, el correo
            inscripcion.postulante.usuario.telefono or inscripcion.postulante.usuario.email or '—',
            inscripcion.carrera.nombre if inscripcion.carrera else 'Sin asignar',
            inscripcion.numero_postulacion,
            inscripcion.convocatoria.nombre,
            inscripcion.fecha_inscripcion.strftime('%d/%m/%Y'),
            'Completo' if inscripcion.proceso_completo else 'Incompleto',
            f'{docs_subidos}/{docs_totales}',
        ]
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=col, value=valor)
            celda.font = Font(name='Calibri', size=10.5, color=negro)
            celda.border = borde_celda
            celda.alignment = Alignment(
                horizontal='left' if col in (2, 4, 5) else 'center',
                vertical='center',
                wrap_text=(col in (2, 4, 5)),
            )
        fila += 1

    if fila == fila_encabezado + 1:
        ws.merge_cells(f'A{fila}:{ultima_letra}{fila}')
        vacio = ws.cell(row=fila, column=1, value='No hay postulantes registrados con estos filtros.')
        vacio.font = Font(name='Calibri', size=10.5, italic=True, color='FF757575')
        vacio.alignment = Alignment(horizontal='center', vertical='center')

    ws.freeze_panes = f'A{fila_encabezado + 1}'
    # Al imprimir el Excel: hoja horizontal y todas las columnas en una página de ancho
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    nombre_archivo = f'postulantes_istam_{timezone.localtime().strftime("%Y%m%d_%H%M")}.xlsx'
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    wb.save(response)
    return response


@panel_required
def detalle_postulante(request, inscripcion_id):
    inscripcion = get_object_or_404(
        Inscripcion.objects.select_related('postulante', 'postulante__usuario', 'carrera', 'convocatoria'),
        id=inscripcion_id,
    )

    docs_por_tipo = {d.tipo: d for d in inscripcion.documentos.all()}
    docs_detalle = [
        {
            'tipo': tipo,
            'etiqueta': etiqueta,
            'documento': docs_por_tipo.get(tipo),
            'es_imagen': bool(docs_por_tipo.get(tipo)) and docs_por_tipo[tipo].archivo.name.lower().endswith(
                ('.jpg', '.jpeg', '.png', '.gif', '.webp')
            ),
        }
        for tipo, etiqueta in Documento.TIPO_CHOICES
    ]
    docs_validados = sum(1 for d in docs_detalle if d['documento'] and d['documento'].estado == 'validado')
    todos_validados = docs_validados == len(docs_detalle)

    resultado = Resultado.objects.filter(inscripcion=inscripcion).select_related('examen').first()

    return render(request, 'panel/detalle.html', {
        'inscripcion': inscripcion,
        'docs_detalle': docs_detalle,
        'docs_validados': docs_validados,
        'total_docs': len(docs_detalle),
        'todos_validados': todos_validados,
        'resultado': resultado,
    })


@panel_required
def validar_documento(request, documento_id):
    documento = get_object_or_404(Documento, id=documento_id)

    if request.method != 'POST':
        return redirect('panel:detalle_postulante', inscripcion_id=documento.inscripcion_id)

    accion = request.POST.get('accion')
    observaciones = request.POST.get('observaciones', '').strip()

    if accion not in ('aprobar', 'rechazar'):
        messages.error(request, 'Acción no reconocida.')
        return redirect('panel:detalle_postulante', inscripcion_id=documento.inscripcion_id)

    documento.estado = 'validado' if accion == 'aprobar' else 'rechazado'
    documento.observaciones = observaciones
    documento.fecha_revision = timezone.now()
    documento.revisado_por = request.user
    documento.save()

    if accion == 'rechazar' and documento.inscripcion.estado == 'aprobada':
        documento.inscripcion.estado = 'docs_pendientes'
        documento.inscripcion.save(update_fields=['estado'])
        documento.inscripcion.actualizar_estado_por_documentos()
    notificar_documento(documento)

    if accion == 'aprobar':
        messages.success(request, f'{documento.get_tipo_display()} de {documento.inscripcion.postulante.nombre_completo} marcado como validado.')
    else:
        messages.warning(request, f'{documento.get_tipo_display()} de {documento.inscripcion.postulante.nombre_completo} rechazado. El postulante deberá volver a subirlo.')

    destino = request.POST.get('next', '')
    if destino and url_has_allowed_host_and_scheme(destino, allowed_hosts={request.get_host()}):
        return redirect(destino)
    return redirect('panel:detalle_postulante', inscripcion_id=documento.inscripcion_id)


@panel_required
def decidir_inscripcion(request, inscripcion_id):
    inscripcion = get_object_or_404(Inscripcion, id=inscripcion_id)

    if request.method != 'POST':
        return redirect('panel:detalle_postulante', inscripcion_id=inscripcion.id)

    accion = request.POST.get('accion')
    motivo = (request.POST.get('motivo') or '').strip()[:800]
    volver = redirect('panel:detalle_postulante', inscripcion_id=inscripcion.id)
    etiquetas = dict(Documento.TIPO_CHOICES)

    if accion == 'aprobar':
        # Validar proceso: solo cuando los 5 documentos están validados
        tipos = list(etiquetas)
        pendientes = inscripcion.documentos.filter(tipo__in=tipos).exclude(estado='validado').count()
        faltantes = len(tipos) - inscripcion.documentos.filter(tipo__in=tipos).count()
        if pendientes or faltantes:
            messages.error(request, 'No puedes validar el proceso: todavía hay documentos sin validar.')
            return volver
        inscripcion.estado = 'aprobada'
        inscripcion.save(update_fields=['estado'])
        notificar_decision_inscripcion(inscripcion)
        messages.success(request, f'Proceso de {inscripcion.postulante.nombre_completo} validado. Se le avisó por correo.')

    elif accion == 'devolver':
        # Rechazar proceso pidiendo corrección: los documentos marcados quedan
        # "rechazados" con el motivo y el postulante puede volver a subirlos.
        tipos = [t for t in request.POST.getlist('documentos') if t in etiquetas]
        if not tipos:
            messages.error(request, 'Marca al menos un documento que el postulante debe volver a subir.')
            return volver
        if not motivo:
            messages.error(request, 'Escribe el motivo para que el postulante sepa qué corregir.')
            return volver
        ahora = timezone.now()
        with transaction.atomic():
            for doc in inscripcion.documentos.filter(tipo__in=tipos):
                doc.estado = 'rechazado'
                doc.observaciones = motivo
                doc.fecha_revision = ahora
                doc.revisado_por = request.user
                doc.save(update_fields=['estado', 'observaciones', 'fecha_revision', 'revisado_por'])
            nombres = ', '.join(etiquetas[t] for t in tipos)
            inscripcion.estado = 'docs_pendientes'
            inscripcion.comentario_revision = f'Debes volver a subir: {nombres}. Motivo: {motivo}'
            inscripcion.fecha_comentario = ahora
            inscripcion.save(update_fields=['estado', 'comentario_revision', 'fecha_comentario'])
        notificar_comentario(inscripcion)
        messages.warning(
            request,
            f'Proceso devuelto a {inscripcion.postulante.nombre_completo}: debe volver a subir {nombres}. '
            'Se le avisó por correo. Cuando los suba, valídalos y luego pulsa "Validar proceso".'
        )

    elif accion == 'rechazar':
        # Rechazo definitivo: ya no puede continuar en este proceso
        if not motivo:
            messages.error(request, 'Escribe el motivo del rechazo.')
            return volver
        inscripcion.estado = 'rechazada'
        inscripcion.comentario_revision = f'Postulación rechazada. Motivo: {motivo}'
        inscripcion.fecha_comentario = timezone.now()
        inscripcion.save(update_fields=['estado', 'comentario_revision', 'fecha_comentario'])
        notificar_decision_inscripcion(inscripcion)
        messages.warning(request, f'Proceso de {inscripcion.postulante.nombre_completo} rechazado definitivamente.')

    else:
        messages.error(request, 'Acción no reconocida.')

    return redirect('panel:detalle_postulante', inscripcion_id=inscripcion.id)


# ============================================================
# Carreras — CRUD propio del panel (ya no pasa por /admin/)
# ============================================================

@panel_required
def carreras_lista(request):
    carreras = paginar(request, Carrera.objects.all().order_by('nombre'))
    return render(request, 'panel/carreras_lista.html', {'carreras': carreras})


@panel_required
def carrera_crear(request):
    if request.method == 'POST':
        form = CarreraForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            messages.success(request, 'Carrera creada correctamente.')
            return redirect('panel:carreras_lista')
    else:
        form = CarreraForm()
    return render(request, 'panel/carrera_form.html', {'form': form, 'es_nueva': True})


@panel_required
def carrera_editar(request, carrera_id):
    carrera = get_object_or_404(Carrera, id=carrera_id)
    if request.method == 'POST':
        form = CarreraForm(request.POST, request.FILES, instance=carrera)
        if form.is_valid():
            form.save()
            messages.success(request, 'Carrera actualizada correctamente.')
            return redirect('panel:carreras_lista')
    else:
        form = CarreraForm(instance=carrera)
    return render(request, 'panel/carrera_form.html', {'form': form, 'carrera': carrera, 'es_nueva': False})


@panel_required
def carrera_cambiar_estado(request, carrera_id):
    carrera = get_object_or_404(Carrera, id=carrera_id)
    if request.method == 'POST':
        carrera.activa = not carrera.activa
        carrera.save(update_fields=['activa'])
        if carrera.activa:
            messages.success(request, f'La carrera "{carrera.nombre}" fue activada.')
        else:
            messages.success(request, f'La carrera "{carrera.nombre}" fue pausada.')
    return redirect('panel:carreras_lista')


@panel_required
def carreras_eliminar(request):
    if request.method != 'POST':
        return redirect('panel:carreras_lista')

    ids = [i for i in request.POST.getlist('carreras') if i.isdigit()]
    if not ids:
        messages.error(request, 'No marcaste ninguna carrera.')
        return redirect('panel:carreras_lista')

    eliminadas, con_postulantes = [], []
    for carrera in Carrera.objects.filter(id__in=ids):
        if Inscripcion.objects.filter(carrera=carrera).exists():
            con_postulantes.append(carrera.nombre)
        else:
            carrera.delete()
            eliminadas.append(carrera.nombre)

    if eliminadas:
        messages.success(request, f'Se eliminó: {", ".join(eliminadas)}.')
    if con_postulantes:
        messages.warning(
            request,
            f'No se pudo eliminar {", ".join(con_postulantes)} porque tiene postulantes inscritos. '
            'Puedes pausarla para que no se ofrezca más.'
        )
    return redirect('panel:carreras_lista')


# ============================================================
# Procesos de admisión (Convocatorias) — con cupos por carrera
# ============================================================

@panel_required
def convocatorias_lista(request):
    convocatorias = list(
        Convocatoria.objects.annotate(total_carreras=Count('cupocarrera', distinct=True))
        .order_by('-fecha_inicio')
    )
    hoy = date.today()

    for conv in convocatorias:
        if conv.fecha_fin < hoy:
            conv.etiqueta, conv.etiqueta_clase = 'Proceso finalizado', 'etiqueta-cerrado'
        elif conv.fecha_inicio > hoy:
            conv.etiqueta, conv.etiqueta_clase = 'Nuevo proceso', 'etiqueta-nuevo'
        else:
            conv.etiqueta, conv.etiqueta_clase = 'Proceso actual', 'etiqueta-actual'

    return render(request, 'panel/convocatorias_lista.html', {'convocatorias': paginar(request, convocatorias)})


@panel_required
def convocatoria_crear(request):
    # Si se llega desde el detalle de un proceso anterior ("Nuevo proceso"),
    # al guardar se vuelve a ese proceso para reintegrar a sus estudiantes.
    origen_id = request.POST.get('origen') or request.GET.get('origen')
    origen = Convocatoria.objects.filter(id=origen_id).first() if origen_id and str(origen_id).isdigit() else None

    if request.method == 'POST':
        form = ConvocatoriaForm(request.POST)
        formset = CupoCarreraFormSet(request.POST, instance=form.instance)
        if form.is_valid() and formset.is_valid():
            convocatoria = form.save()
            formset.instance = convocatoria
            formset.save()
            messages.success(request, 'Proceso de admisión creado correctamente.')
            if origen:
                messages.info(request, f'Ahora selecciona los estudiantes de "{origen.nombre}" que quieres pasar a "{convocatoria.nombre}".')
                return redirect(f"{reverse('panel:convocatoria_detalle', args=[origen.id])}?destino={convocatoria.id}")
            return redirect('panel:convocatorias_lista')
    else:
        form = ConvocatoriaForm()
        formset = CupoCarreraFormSet()
    return render(request, 'panel/convocatoria_form.html', {
        'form': form, 'formset': formset, 'es_nueva': True, 'origen': origen,
    })


def _resumen_convocatoria(convocatoria):
    """Postulantes del proceso según su avance, en total y por carrera:
    completo (subió los 5 documentos) o incompleto (le falta alguno,
    debe volver a subir alguno o no ha subido nada)."""
    inscripciones = list(
        Inscripcion.objects.filter(convocatoria=convocatoria).prefetch_related('documentos')
    )
    por_carrera = []
    for cupo in convocatoria.cupocarrera_set.select_related('carrera').order_by('carrera__nombre'):
        de_carrera = [i for i in inscripciones if i.carrera_id == cupo.carrera_id]
        completos = sum(1 for i in de_carrera if i.proceso_completo)
        por_carrera.append({
            'carrera': cupo.carrera,
            'cupos': cupo.cupos,
            'postulantes': len(de_carrera),
            'completos': completos,
            'incompletos': len(de_carrera) - completos,
        })
    completos = sum(1 for i in inscripciones if i.proceso_completo)
    return {
        'postulantes': len(inscripciones),
        'completos': completos,
        'incompletos': len(inscripciones) - completos,
        'vacios': sum(1 for i in inscripciones if i.avance == 'vacio'),
        'por_carrera': por_carrera,
    }


def _destinos_reintegro(convocatoria):
    """Procesos a los que se puede reintegrar: todos menos este y que no
    hayan terminado (actual o nuevo). Primero los activos."""
    return (
        Convocatoria.objects.exclude(id=convocatoria.id)
        .filter(fecha_fin__gte=date.today())
        .order_by('-activa', 'fecha_inicio')
    )


@panel_required
def convocatoria_detalle(request, convocatoria_id):
    """Detalle de un proceso: cuántos lo completaron y cuántos no, con la
    lista de los que no lo completaron para reintegrarlos a un proceso
    abierto (actual o nuevo) o eliminarlos de este proceso."""
    convocatoria = get_object_or_404(Convocatoria, id=convocatoria_id)
    resumen = _resumen_convocatoria(convocatoria)
    hoy = date.today()
    finalizado = convocatoria.fecha_fin < hoy

    destinos = _destinos_reintegro(convocatoria)
    destino = None
    destino_id = request.GET.get('destino', '')
    if destino_id.isdigit():
        destino = destinos.filter(id=destino_id).first()
    if destino is None:
        destino = destinos.first()

    filtro_carrera = request.GET.get('carrera', '')
    ver = request.GET.get('ver', 'pendientes')
    if ver not in ('pendientes', 'reintegrados', 'todos'):
        ver = 'pendientes'
    if not finalizado:
        ver = 'todos'   # en un proceso abierto todavía no se reintegra

    consulta = (
        Inscripcion.objects.filter(convocatoria=convocatoria)
        .select_related('postulante', 'postulante__usuario', 'carrera')
        .prefetch_related('documentos')
        .order_by('postulante__apellidos', 'postulante__nombres')
    )
    if filtro_carrera.isdigit():
        consulta = consulta.filter(carrera_id=filtro_carrera)
    # "No completaron" = les falta algún documento
    quedados = [i for i in consulta if not i.proceso_completo]

    ya_en_destino, carreras_destino = set(), set()
    if destino:
        ya_en_destino = set(
            Inscripcion.objects.filter(convocatoria=destino).values_list('postulante_id', flat=True)
        )
        carreras_destino = set(destino.cupocarrera_set.values_list('carrera_id', flat=True))

    for inscripcion in quedados:
        inscripcion.ya_reintegrado = inscripcion.postulante_id in ya_en_destino
        inscripcion.carrera_en_destino = inscripcion.carrera_id in carreras_destino

    total_pendientes = sum(1 for i in quedados if not i.ya_reintegrado)
    total_reintegrados = len(quedados) - total_pendientes
    if ver == 'pendientes':
        lista = [i for i in quedados if not i.ya_reintegrado]
    elif ver == 'reintegrados':
        lista = [i for i in quedados if i.ya_reintegrado]
    else:
        lista = quedados

    if finalizado:
        estado_proceso = {
            'clase': 'fin',
            'titulo': 'Proceso finalizado',
            'texto': f'Cerró el {convocatoria.fecha_fin:%d/%m/%Y}',
            'numero': (hoy - convocatoria.fecha_fin).days,
            'unidad': 'días desde el cierre',
        }
    elif convocatoria.fecha_inicio > hoy:
        estado_proceso = {
            'clase': 'nuevo',
            'titulo': 'Nuevo proceso',
            'texto': f'Abre las inscripciones el {convocatoria.fecha_inicio:%d/%m/%Y}',
            'numero': (convocatoria.fecha_inicio - hoy).days,
            'unidad': 'días para comenzar',
        }
    else:
        estado_proceso = {
            'clase': 'actual',
            'titulo': 'Proceso actual',
            'texto': f'Inscripciones abiertas hasta el {convocatoria.fecha_fin:%d/%m/%Y}',
            'numero': (hoy - convocatoria.fecha_inicio).days + 1,
            'unidad': f'de {(convocatoria.fecha_fin - convocatoria.fecha_inicio).days + 1} días',
        }
    return render(request, 'panel/convocatoria_detalle.html', {
        'convocatoria': convocatoria,
        'resumen': resumen,
        'destinos': destinos,
        'destino': destino,
        'quedados': paginar(request, lista),
        'ver': ver,
        'total_no_completaron': len(quedados),
        'total_pendientes': total_pendientes,
        'total_reintegrados': total_reintegrados,
        'pendientes_reintegrar': total_pendientes,
        'carreras': [f['carrera'] for f in resumen['por_carrera']],
        'filtro_carrera': filtro_carrera,
        'finalizado': finalizado,
        'total_cupos': sum(f['cupos'] for f in resumen['por_carrera']),
        'es_proceso_actual': (_convocatoria_activa() or Convocatoria()).id == convocatoria.id,
        'estado_proceso': estado_proceso,
        'motivos_reintegro': MOTIVOS_REINTEGRO,
    })


# Motivos que el personal puede elegir al reintegrar. Se envían por correo
# al postulante y quedan como "Comentario de admisión" en su nueva postulación.
MOTIVOS_REINTEGRO = [
    ('auto', 'Automático (según lo que le faltó)'),
    ('documentos', 'No completó sus documentos'),
    ('evaluacion', 'No rindió la evaluación'),
    ('cupo', 'No alcanzó cupo en la carrera'),
    ('solicitud', 'Lo solicitó el postulante'),
    ('otro', 'Otro motivo (escribir mensaje)'),
]


def _concepto_reintegro(motivo, anterior, origen):
    """Texto del motivo que verá el postulante en el correo."""
    if motivo == 'documentos':
        return f'No completaste la entrega de tus documentos en el proceso "{origen.nombre}".'
    if motivo == 'evaluacion':
        return f'No rendiste la evaluación de admisión en el proceso "{origen.nombre}".'
    if motivo == 'cupo':
        return f'No alcanzaste un cupo en la carrera {anterior.carrera.nombre} en el proceso "{origen.nombre}".'
    if motivo == 'solicitud':
        return 'Solicitaste participar nuevamente en el proceso de admisión.'
    if motivo == 'otro':
        return ''
    # Automático: se arma con lo que realmente le faltó en el proceso anterior
    faltan = anterior.documentos_faltantes
    if anterior.avance == 'vacio':
        return f'No subiste ningún documento en el proceso "{origen.nombre}".'
    partes = []
    no_subidos = [d['etiqueta'] for d in faltan if d['situacion'] == 'falta']
    rechazados = [d['etiqueta'] for d in faltan if d['situacion'] == 'rechazado']
    if no_subidos:
        partes.append('te faltó subir: ' + ', '.join(no_subidos))
    if rechazados:
        partes.append('debías volver a subir: ' + ', '.join(rechazados))
    return f'En el proceso "{origen.nombre}" ' + '; '.join(partes) + '.'


@panel_required
def reintegrar_postulantes(request, convocatoria_id):
    """Pasa a los estudiantes que se quedaron en un proceso anterior al
    proceso elegido (actual o nuevo). Crea una inscripción nueva en el
    proceso destino con la misma carrera. La inscripción del proceso
    anterior se conserva como historial."""
    origen = get_object_or_404(Convocatoria, id=convocatoria_id)
    if request.method != 'POST':
        return redirect('panel:convocatoria_detalle', convocatoria_id=origen.id)

    destino_id = request.POST.get('destino', '')
    destino = _destinos_reintegro(origen).filter(id=destino_id).first() if destino_id.isdigit() else None
    if destino is None:
        messages.error(request, 'Elige un proceso abierto (actual o nuevo) al que quieras pasar a los estudiantes.')
        return redirect('panel:convocatoria_detalle', convocatoria_id=origen.id)

    ids = [i for i in request.POST.getlist('inscripciones') if i.isdigit()]
    if not ids:
        messages.error(request, 'No seleccionaste ningún estudiante.')
        return redirect(f"{reverse('panel:convocatoria_detalle', args=[origen.id])}?destino={destino.id}")

    carreras_destino = set(destino.cupocarrera_set.values_list('carrera_id', flat=True))
    ya_en_destino = set(Inscripcion.objects.filter(convocatoria=destino).values_list('postulante_id', flat=True))

    motivo = request.POST.get('motivo', 'auto')
    if motivo not in dict(MOTIVOS_REINTEGRO):
        motivo = 'auto'
    mensaje = (request.POST.get('mensaje') or '').strip()[:500]
    volver = f"{reverse('panel:convocatoria_detalle', args=[origen.id])}?destino={destino.id}"
    if motivo == 'otro' and not mensaje:
        messages.error(request, 'Elegiste "Otro motivo": escribe el mensaje que recibirá el estudiante.')
        return redirect(volver)
    enlace = request.build_absolute_uri(reverse('mis_inscripciones'))

    reintegrados, repetidos, sin_carrera = 0, 0, []
    por_avisar = []   # (nueva inscripción, concepto) — se envían al terminar de guardar
    seleccion = [
        i for i in Inscripcion.objects.filter(id__in=ids, convocatoria=origen)
        .select_related('postulante', 'postulante__usuario', 'carrera')
        .prefetch_related('documentos')
        if not i.proceso_completo
    ]
    with transaction.atomic():
        for anterior in seleccion:
            if anterior.postulante_id in ya_en_destino:
                repetidos += 1
                continue
            if anterior.carrera_id not in carreras_destino:
                sin_carrera.append(f'{anterior.postulante.nombre_completo} ({anterior.carrera.nombre})')
                continue
            # CORREGIDO: antes se creaba la inscripción dos veces (duplicados)
            nueva = Inscripcion.objects.create(
                postulante=anterior.postulante,
                convocatoria=destino,
                carrera=anterior.carrera,
            )
            nueva.actualizar_estado_por_documentos()
            concepto = _concepto_reintegro(motivo, anterior, origen)
            # Queda visible para el postulante en "Mis inscripciones"
            nueva.comentario_revision = ' '.join(t for t in [
                f'Reintegrado desde "{origen.nombre}".', concepto, mensaje
            ] if t)
            nueva.fecha_comentario = timezone.now()
            nueva.save(update_fields=['comentario_revision', 'fecha_comentario'])
            por_avisar.append((nueva, concepto))
            ya_en_destino.add(anterior.postulante_id)
            reintegrados += 1

    # El correo se envía siempre, automáticamente, a cada reintegrado.
    avisados, sin_correo = 0, []
    for nueva, concepto in por_avisar:
        if notificar_reintegro(nueva, origen, concepto=concepto, mensaje=mensaje, enlace=enlace):
            avisados += 1
        else:
            sin_correo.append(nueva.postulante.nombre_completo)

    if reintegrados:
        messages.success(
            request,
            f'{reintegrados} estudiante(s) reintegrado(s) a "{destino.nombre}". '
            f'Se envió el aviso por correo a {avisados}.'
        )
    if sin_correo:
        messages.warning(
            request,
            f'{len(sin_correo)} estudiante(s) no tienen correo registrado y no recibieron el aviso: '
            f'{", ".join(sin_correo)}. Verán el motivo al ingresar a "Mis inscripciones".'
        )
    if repetidos:
        messages.info(request, f'{repetidos} estudiante(s) ya estaban inscritos en "{destino.nombre}" y no se duplicaron.')
    if sin_carrera:
        messages.warning(
            request,
            f'No se pudo reintegrar a {len(sin_carrera)} estudiante(s) porque su carrera no tiene cupos en '
            f'"{destino.nombre}": {", ".join(sin_carrera)}. Agrega esa carrera al proceso y vuelve a intentarlo.'
        )
    return redirect(volver)


@panel_required
def eliminar_postulantes(request, convocatoria_id):
    """Elimina de este proceso las inscripciones marcadas (y sus documentos).
    La cuenta y los datos personales del postulante se conservan."""
    from django.db.models import ProtectedError

    convocatoria = get_object_or_404(Convocatoria, id=convocatoria_id)
    volver = reverse('panel:convocatoria_detalle', args=[convocatoria.id])
    destino_id = request.POST.get('destino', '')
    if destino_id.isdigit():
        volver = f'{volver}?destino={destino_id}'

    if request.method != 'POST':
        return redirect(volver)

    ids = [i for i in request.POST.getlist('inscripciones') if i.isdigit()]
    if not ids:
        messages.error(request, 'No seleccionaste ningún postulante.')
        return redirect(volver)

    seleccion = Inscripcion.objects.filter(id__in=ids, convocatoria=convocatoria)
    total = seleccion.count()
    try:
        with transaction.atomic():
            for inscripcion in seleccion:
                for doc in inscripcion.documentos.all():
                    if doc.archivo:
                        doc.archivo.delete(save=False)
                inscripcion.delete()
    except ProtectedError:
        messages.error(
            request,
            'No se pudo eliminar: alguno de los postulantes ya tiene resultados u otros registros asociados.'
        )
        return redirect(volver)

    messages.success(request, f'{total} postulante(s) eliminado(s) de "{convocatoria.nombre}".')
    return redirect(volver)


@panel_required
def convocatoria_editar(request, convocatoria_id):
    convocatoria = get_object_or_404(Convocatoria, id=convocatoria_id)
    if request.method == 'POST':
        form = ConvocatoriaForm(request.POST, instance=convocatoria)
        formset = CupoCarreraFormSet(request.POST, instance=convocatoria)
        if form.is_valid() and formset.is_valid():
            form.save()
            formset.save()
            messages.success(request, 'Proceso de admisión actualizado correctamente.')
            return redirect('panel:convocatorias_lista')
    else:
        form = ConvocatoriaForm(instance=convocatoria)
        formset = CupoCarreraFormSet(instance=convocatoria)
    return render(request, 'panel/convocatoria_form.html', {
        'form': form, 'formset': formset, 'convocatoria': convocatoria, 'es_nueva': False,
    })


@panel_required
def convocatoria_cambiar_estado(request, convocatoria_id):
    convocatoria = get_object_or_404(Convocatoria, id=convocatoria_id)
    if request.method == 'POST':
        convocatoria.activa = not convocatoria.activa
        convocatoria.save(update_fields=['activa'])
        if convocatoria.activa:
            messages.success(request, f'El proceso "{convocatoria.nombre}" fue activado. Ya aparece en Carreras.')
        else:
            messages.success(request, f'El proceso "{convocatoria.nombre}" fue pausado. Ya no aparece en Carreras.')
    return redirect('panel:convocatorias_lista')


@panel_required
def convocatorias_eliminar(request):
    """NUEVO: elimina los procesos marcados con casilla en la lista.
    No elimina los que ya tienen postulantes inscritos (se sugiere pausarlos)."""
    if request.method != 'POST':
        return redirect('panel:convocatorias_lista')

    ids = [i for i in request.POST.getlist('convocatorias') if i.isdigit()]
    if not ids:
        messages.error(request, 'No marcaste ningún proceso.')
        return redirect('panel:convocatorias_lista')

    eliminados, con_postulantes = [], []
    for conv in Convocatoria.objects.filter(id__in=ids):
        if Inscripcion.objects.filter(convocatoria=conv).exists():
            con_postulantes.append(conv.nombre)
        else:
            conv.delete()
            eliminados.append(conv.nombre)

    if eliminados:
        messages.success(request, f'Se eliminó: {", ".join(eliminados)}.')
    if con_postulantes:
        messages.warning(
            request,
            f'No se pudo eliminar {", ".join(con_postulantes)} porque tiene postulantes inscritos. '
            'Puedes pausarlo para que no se ofrezca más.'
        )
    return redirect('panel:convocatorias_lista')


# ============================================================
# Exámenes
# ============================================================

@panel_required
def examenes_lista(request):
    examenes = paginar(request, Examen.objects.all().order_by('-fecha', '-hora'))
    return render(request, 'panel/examenes_lista.html', {'examenes': examenes})


@panel_required
def examen_crear(request):
    if request.method == 'POST':
        form = ExamenForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Examen creado correctamente.')
            return redirect('panel:examenes_lista')
    else:
        form = ExamenForm()
    return render(request, 'panel/examen_form.html', {'form': form, 'es_nuevo': True})


@panel_required
def examen_editar(request, examen_id):
    examen = get_object_or_404(Examen, id=examen_id)
    if request.method == 'POST':
        form = ExamenForm(request.POST, instance=examen)
        if form.is_valid():
            form.save()
            messages.success(request, 'Examen actualizado correctamente.')
            return redirect('panel:examenes_lista')
    else:
        form = ExamenForm(instance=examen)
    return render(request, 'panel/examen_form.html', {'form': form, 'examen': examen, 'es_nuevo': False})


# ============================================================
# Reportes (solo lectura, no pasa por /admin/)
# ============================================================

def _historial_excel(historial, proceso_elegido, carrera_elegida):
    """Excel del historial de postulantes de Reportes, con los mismos
    filtros (proceso y carrera) que se ven en pantalla."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    columnas = [
        ('N°', 6),
        ('Postulante', 32),
        ('Cédula', 15),
        ('Contacto', 30),
        ('Carrera', 32),
        ('Proceso de admisión', 30),
        ('N.° de postulación', 20),
        ('Fecha', 13),
        ('Estado', 22),
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Historial'

    verde_fuerte = 'FF8BC34A'
    verde_claro = 'FFE8F5E9'
    borde_fino = Side(style='thin', color='FFBDBDBD')
    borde = Border(left=borde_fino, right=borde_fino, top=borde_fino, bottom=borde_fino)
    ultima = get_column_letter(len(columnas))

    ws.merge_cells(f'A1:{ultima}2')
    ws['A1'].value = 'INSTITUTO SUPERIOR TECNOLÓGICO AMAZÓNICO'
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='FFFFFFFF')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    for fila_t in (1, 2):
        for col in range(1, len(columnas) + 1):
            ws.cell(row=fila_t, column=col).fill = PatternFill('solid', fgColor=verde_fuerte)

    subtitulo = ' · '.join([
        'Historial de postulantes',
        f'Proceso: {proceso_elegido.nombre if proceso_elegido else "Todos los procesos"}',
        f'Carrera: {carrera_elegida.nombre if carrera_elegida else "Todas"}',
        f'Generado el {timezone.localtime().strftime("%d/%m/%Y %H:%M")}',
    ])
    ws.merge_cells(f'A3:{ultima}3')
    ws['A3'].value = subtitulo
    ws['A3'].font = Font(name='Calibri', size=10, italic=True, color='FF33691E')
    ws['A3'].alignment = Alignment(horizontal='center', vertical='center')

    for col, (etiqueta, ancho) in enumerate(columnas, start=1):
        celda = ws.cell(row=4, column=col, value=etiqueta)
        celda.font = Font(name='Calibri', size=11, bold=True)
        celda.fill = PatternFill('solid', fgColor=verde_claro)
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        celda.border = borde
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.row_dimensions[4].height = 26

    fila = 5
    for numero, ins in enumerate(historial, start=1):
        usuario = ins.postulante.usuario
        contacto = usuario.telefono or usuario.email or '—'
        valores = [
            numero,
            f'{ins.postulante.apellidos} {ins.postulante.nombres}'.strip(),
            usuario.cedula or '—',
            contacto,
            ins.carrera.nombre if ins.carrera else 'Sin asignar',
            ins.convocatoria.nombre,
            ins.numero_postulacion,
            ins.fecha_inscripcion.strftime('%d/%m/%Y'),
            ins.get_estado_display(),
        ]
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=col, value=valor)
            celda.font = Font(name='Calibri', size=10.5)
            celda.border = borde
            celda.alignment = Alignment(horizontal='left' if col in (2, 4, 5, 6) else 'center', vertical='center', wrap_text=True)
        fila += 1

    ws.cell(row=fila + 1, column=1, value=f'Total: {fila - 5} postulantes').font = Font(name='Calibri', size=10, bold=True)
    ws.freeze_panes = 'A5'
    # Al imprimir el Excel: hoja horizontal y todas las columnas en una página de ancho
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    respuesta = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    nombre = f'historial_postulantes_{timezone.localtime().strftime("%Y%m%d_%H%M")}.xlsx'
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre}"'
    wb.save(respuesta)
    return respuesta


@panel_required
def reportes(request):
    convocatoria = _convocatoria_activa()

    por_estado = list(
        Inscripcion.objects.values('estado').annotate(total=Count('id')).order_by('estado')
    )
    estados_dict = dict(Inscripcion.ESTADO_CHOICES)
    for fila in por_estado:
        fila['etiqueta'] = estados_dict.get(fila['estado'], fila['estado'])

    por_carrera = list(
        Inscripcion.objects.values('carrera__nombre').annotate(total=Count('id')).order_by('-total')
    )

    documentos_por_estado = list(
        Documento.objects.values('estado').annotate(total=Count('id')).order_by('estado')
    )
    docs_estados_dict = dict(Documento.ESTADO_CHOICES)
    for fila in documentos_por_estado:
        fila['etiqueta'] = docs_estados_dict.get(fila['estado'], fila['estado'])

    # Historial de postulantes para imprimir: de todos los procesos
    # (anteriores y actual) o de uno en particular, y opcionalmente por carrera.
    historial_proceso = request.GET.get('historial_proceso', '')
    historial_carrera = request.GET.get('historial_carrera', '')
    historial = Inscripcion.objects.select_related(
        'postulante', 'postulante__usuario', 'carrera', 'convocatoria'
    ).order_by('-convocatoria__fecha_inicio', 'carrera__nombre', 'postulante__apellidos', 'postulante__nombres')
    proceso_elegido = None
    carrera_elegida = None
    if historial_proceso.isdigit():
        historial = historial.filter(convocatoria_id=historial_proceso)
        proceso_elegido = Convocatoria.objects.filter(pk=historial_proceso).first()
    if historial_carrera.isdigit():
        historial = historial.filter(carrera_id=historial_carrera)
        carrera_elegida = Carrera.objects.filter(pk=historial_carrera).first()

    # Botón "Exportar a Excel": misma página, con ?exportar=excel
    if request.GET.get('exportar') == 'excel':
        return _historial_excel(historial, proceso_elegido, carrera_elegida)

    return render(request, 'panel/reportes.html', {
        'convocatoria': convocatoria,
        'por_estado': por_estado,
        'por_carrera': por_carrera,
        'documentos_por_estado': documentos_por_estado,
        'total_inscripciones': Inscripcion.objects.count(),
        'historial': historial,
        'historial_pagina': paginar(request, historial),
        'historial_convocatorias': Convocatoria.objects.order_by('-fecha_inicio'),
        'historial_carreras': Carrera.objects.order_by('nombre'),
        'historial_proceso': historial_proceso,
        'historial_carrera': historial_carrera,
        'proceso_elegido': proceso_elegido,
        'carrera_elegida': carrera_elegida,
        'fecha_impresion': timezone.localtime(),
    })


# ============================================================
# Vista previa de documentos Word (.docx) dentro del visor
# ============================================================

@xframe_options_sameorigin
@panel_required
def vista_documento(request, documento_id):
    """Página que se muestra dentro del visor para documentos Word."""
    documento = get_object_or_404(Documento, id=documento_id)
    return render(request, 'panel/vista_documento.html', {'documento': documento})


def _nombre_archivo(texto):
    import re
    import unicodedata
    texto = unicodedata.normalize('NFD', texto).encode('ascii', 'ignore').decode()
    return re.sub(r'[^A-Za-z0-9]+', '_', texto).strip('_')


@panel_required
def descargar_documentos(request, inscripcion_id):
    """Descarga en un solo .zip todos los documentos que subió el postulante."""
    import io
    import os
    import zipfile

    inscripcion = get_object_or_404(Inscripcion.objects.select_related('postulante'), id=inscripcion_id)

    nombre_postulante = _nombre_archivo(inscripcion.postulante.nombre_completo)
    memoria = io.BytesIO()
    agregados = 0
    with zipfile.ZipFile(memoria, 'w', zipfile.ZIP_DEFLATED) as z:
        for doc in inscripcion.documentos.all():
            try:
                with doc.archivo.open('rb') as f:
                    extension = os.path.splitext(doc.archivo.name)[1].lower()
                    z.writestr(f'{_nombre_archivo(doc.get_tipo_display())}_{nombre_postulante}{extension}', f.read())
                    agregados += 1
            except (FileNotFoundError, OSError):
                continue

    if not agregados:
        messages.error(request, 'Este postulante no tiene documentos para descargar.')
        return redirect('panel:detalle_postulante', inscripcion_id=inscripcion.id)

    response = HttpResponse(memoria.getvalue(), content_type='application/zip')
    response['Content-Disposition'] = f'attachment; filename="Documentos_{nombre_postulante}_{inscripcion.numero_postulacion}.zip"'
    return response


@panel_required
def descargar_documento(request, documento_id):
    """Descarga un documento con un nombre claro: Tipo_Nombre_del_postulante.ext"""
    import os
    from django.http import FileResponse, Http404

    documento = get_object_or_404(Documento.objects.select_related('inscripcion__postulante'), id=documento_id)
    try:
        archivo = documento.archivo.open('rb')
    except (FileNotFoundError, OSError):
        raise Http404('El archivo no existe en el servidor.')
    extension = os.path.splitext(documento.archivo.name)[1].lower()
    nombre = f'{_nombre_archivo(documento.get_tipo_display())}_{_nombre_archivo(documento.inscripcion.postulante.nombre_completo)}{extension}'
    return FileResponse(archivo, as_attachment=True, filename=nombre)


@panel_required
def config_pasos_guardar(request):
    if request.method == 'POST':
        form = ConfiguracionProcesoForm(request.POST, instance=ConfiguracionProceso.obtener())
        if form.is_valid():
            form.save()
            messages.success(request, 'Se guardó la configuración del proceso del postulante.')
        else:
            errores = ' '.join(str(e) for lista in form.errors.values() for e in lista)
            messages.error(request, f'No se pudo guardar: {errores}')
    return redirect('panel:configuracion')